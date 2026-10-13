from __future__ import annotations

import shutil
import zipfile
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from openpyxl import Workbook
from pypdf import PdfReader

from app.workers.processors import document_convert
from app.workers.processors.document_convert import _detect_format, convert

requires_soffice = pytest.mark.skipif(
    shutil.which("soffice") is None,
    reason="LibreOffice (soffice) nao instalado",
)

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

_DOCX_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" '
    'ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument'
    '.wordprocessingml.document.main+xml"/>'
    "</Types>"
)
_DOCX_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/>'
    "</Relationships>"
)
_DOCX_DOCUMENT = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    "<w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body>"
    "</w:document>"
)
_ODT_MANIFEST = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<manifest:manifest '
    'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
    'manifest:version="1.2">'
    '<manifest:file-entry manifest:full-path="/" '
    'manifest:media-type="application/vnd.oasis.opendocument.text"/>'
    '<manifest:file-entry manifest:full-path="content.xml" '
    'manifest:media-type="text/xml"/>'
    "</manifest:manifest>"
)
_ODT_CONTENT = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    "<office:document-content "
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'office:version="1.2">'
    "<office:body><office:text><text:p>{text}</text:p></office:text></office:body>"
    "</office:document-content>"
)


def _zip(entries: dict[str, str], stored_first: str | None = None) -> BytesIO:
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        if stored_first is not None:
            zf.writestr(
                stored_first,
                entries[stored_first],
                compress_type=zipfile.ZIP_STORED,
            )
        for name, content in entries.items():
            if name != stored_first:
                zf.writestr(name, content)
    buf.seek(0)
    return buf


def _make_docx(text: str) -> BytesIO:
    return _zip({
        "[Content_Types].xml": _DOCX_CONTENT_TYPES,
        "_rels/.rels": _DOCX_RELS,
        "word/document.xml": _DOCX_DOCUMENT.format(text=text),
    })


def _make_odf(mimetype: str, text: str = "X") -> BytesIO:
    return _zip(
        {
            "mimetype": mimetype,
            "META-INF/manifest.xml": _ODT_MANIFEST,
            "content.xml": _ODT_CONTENT.format(text=text),
        },
        stored_first="mimetype",
    )


def _make_odt(text: str) -> BytesIO:
    return _make_odf("application/vnd.oasis.opendocument.text", text)


def _make_xlsx(text: str) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = text
    ws["B1"] = 42
    ws.column_dimensions["A"].width = 30
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _make_rtf(text: str) -> BytesIO:
    rtf = (
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Liberation Serif;}}"
        f"\\f0\\fs24 {text}\\par}}"
    )
    return BytesIO(rtf.encode("ascii"))


def _pdf_text(payload: BytesIO) -> str:
    reader = PdfReader(payload)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


@pytest.mark.parametrize(  # deteccao de formato (roda sem soffice)
    ("factory", "expected"),
    [
        (lambda: _make_rtf("X"), "rtf"),
        (lambda: _make_docx("X"), "docx"),
        (lambda: _make_xlsx("X"), "xlsx"),
        (lambda: _zip({
            "[Content_Types].xml": "<Types/>",
            "ppt/presentation.xml": "<p/>",
        }), "pptx"),
    ],
)
def test_detecta_formatos_texto_e_ooxml(factory, expected):
    assert _detect_format(factory().getvalue()) == expected


@pytest.mark.parametrize(
    ("mimetype", "expected"),
    [
        ("application/vnd.oasis.opendocument.text", "odt"),
        ("application/vnd.oasis.opendocument.spreadsheet", "ods"),
        ("application/vnd.oasis.opendocument.presentation", "odp"),
    ],
)
def test_detecta_opendocument(mimetype, expected):
    assert _detect_format(_make_odf(mimetype).getvalue()) == expected


@pytest.mark.parametrize(
    ("stream_name", "expected"),
    [
        ("WordDocument", "doc"),
        ("Workbook", "xls"),
        ("PowerPoint Document", "ppt"),
    ],
)
def test_detecta_ole2_legado(stream_name, expected):
    data = _OLE2_MAGIC + b"\x00" * 504 + stream_name.encode("utf-16-le")
    assert _detect_format(data) == expected


def test_ole2_doc_com_planilha_embutida_continua_doc():
    data = (
            _OLE2_MAGIC
            + "Workbook".encode("utf-16-le")
            + "WordDocument".encode("utf-16-le")
    )
    assert _detect_format(data) == "doc"


def test_ole2_desconhecido_raises():
    with pytest.raises(ValueError, match="OLE2 nao reconhecido"):
        _detect_format(_OLE2_MAGIC + b"\x00" * 512)


def test_odf_mimetype_nao_suportado_raises():
    data = _make_odf("application/vnd.oasis.opendocument.graphics").getvalue()
    with pytest.raises(ValueError, match="OpenDocument nao suportado"):
        _detect_format(data)


def test_ooxml_sem_diretorio_conhecido_raises():
    data = _zip({"[Content_Types].xml": "<Types/>", "foo/bar.xml": "<x/>"})
    with pytest.raises(ValueError, match="OOXML nao reconhecido"):
        _detect_format(data.getvalue())


def test_zip_generico_raises():
    data = _zip({"leia-me.txt": "nada de office aqui"})
    with pytest.raises(ValueError, match="nem OpenDocument"):
        _detect_format(data.getvalue())


def test_zip_corrompido_raises():
    with pytest.raises(ValueError, match="ZIP invalido"):
        _detect_format(b"PK\x03\x04" + b"\x00" * 100)


@pytest.mark.parametrize(
    "data",
    [
        b"texto puro qualquer",
        b"%PDF-1.7 ja sou pdf",
        b"\x89PNG\r\n\x1a\n" + b"\x00" * 20,
    ],
)
def test_formato_nao_suportado_raises_antes_do_soffice(data):
    with pytest.raises(ValueError, match="formato nao suportado"):
        convert(BytesIO(data), {})


@requires_soffice  # conversao real via LibreOffice
def test_docx_para_pdf_preserva_texto():
    result = convert(_make_docx("Relatorio Trimestral"), {})

    assert "Relatorio Trimestral" in _pdf_text(result.output)


@requires_soffice
def test_odt_para_pdf_preserva_texto():
    result = convert(_make_odt("Documento OpenDocument"), {})

    assert "Documento OpenDocument" in _pdf_text(result.output)


@requires_soffice
def test_rtf_para_pdf_preserva_texto():
    result = convert(_make_rtf("Texto em RTF"), {})

    assert "Texto em RTF" in _pdf_text(result.output)


@requires_soffice
def test_xlsx_para_pdf_preserva_celulas():
    result = convert(_make_xlsx("Planilha Vendas"), {})

    body = _pdf_text(result.output)
    assert "Planilha Vendas" in body
    assert "42" in body


@requires_soffice
def test_content_type_e_extensao():
    result = convert(_make_docx("X"), {})

    assert result.output_content_type == "application/pdf"
    assert result.output_extension == "pdf"
    assert result.output.getvalue().startswith(b"%PDF-")


@requires_soffice
def test_metadata():
    result = convert(_make_docx("X"), {})

    meta = result.result_metadata
    assert meta["source_format"] == "docx"
    assert meta["pages"] == 1
    assert meta["output_size_bytes"] == len(result.output.getvalue())
    assert meta["soffice_duration_seconds"] >= 0


@requires_soffice
def test_docx_com_xml_corrompido_raises():
    data = _zip({
        "[Content_Types].xml": _DOCX_CONTENT_TYPES,
        "_rels/.rels": _DOCX_RELS,
        "word/document.xml": "isto <nao> e xml valido",
    })
    with pytest.raises(ValueError, match="nao conseguiu converter"):
        convert(data, {})


@requires_soffice
def test_timeout_mata_soffice_e_raises(monkeypatch):
    monkeypatch.setattr(document_convert, "DOCUMENT_TIMEOUT_SECONDS", 0.001)

    with pytest.raises(ValueError, match="timeout"):
        convert(_make_docx("X"), {})


@requires_soffice
def test_conversoes_concorrentes_nao_conflitam():
    texts = ["Primeiro documento", "Segundo documento"]

    with ThreadPoolExecutor(max_workers=2) as pool:
        results = list(pool.map(lambda t: convert(_make_docx(t), {}), texts))

    for text, result in zip(texts, results):
        assert text in _pdf_text(result.output)
