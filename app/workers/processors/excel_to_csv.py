from __future__ import annotations

import csv
from datetime import date, datetime, time
from io import BytesIO, StringIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.workers.processors.base import ProcessorResult


def to_csv(
    input_stream: BytesIO,
    parameters: dict[str, Any],
) -> ProcessorResult:
    sheet_param = parameters.get("sheet")
    delimiter = parameters.get("delimiter", ",")
    encoding = parameters.get("encoding", "utf-8")
    include_header = bool(parameters.get("include_header", True))

    if not isinstance(delimiter, str) or len(delimiter) != 1:
        raise ValueError("'delimiter' deve ser uma string de 1 caractere")

    try:
        "".encode(encoding)
    except LookupError as exc:
        raise ValueError(f"encoding invalido: '{encoding}'") from exc

    try:
        wb = load_workbook(input_stream, read_only=True, data_only=True)
    except (InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ValueError(f"arquivo Excel invalido: {exc}") from exc

    try:
        ws = _resolve_sheet(wb, sheet_param)
        sheet_name = ws.title

        text_buf = StringIO(newline="")
        writer = csv.writer(text_buf, delimiter=delimiter, lineterminator="\n")

        columns = 0
        data_rows = 0

        for idx, row in enumerate(ws.iter_rows(values_only=True)):
            if idx == 0 and not include_header:
                continue
            cells = [_normalize(v) for v in row]
            if columns == 0:
                columns = len(cells)
            writer.writerow(cells)
            if include_header and idx == 0:
                continue
            data_rows += 1

        payload = text_buf.getvalue().encode(encoding)
    finally:
        wb.close()

    return ProcessorResult(
        output=BytesIO(payload),
        output_content_type=f"text/csv; charset={encoding}",
        output_extension="csv",
        result_metadata={
            "sheet_name": sheet_name,
            "rows_written": data_rows,
            "columns": columns,
            "encoding": encoding,
            "delimiter": delimiter,
        },
    )


def _resolve_sheet(wb, sheet_param: Any):
    sheets = wb.worksheets
    if not sheets:
        raise ValueError("workbook sem planilhas")

    if sheet_param is None:
        return sheets[0]

    if isinstance(sheet_param, bool):
        raise ValueError("'sheet' deve ser nome (string) ou indice (int)")

    if isinstance(sheet_param, int):
        if sheet_param < 0 or sheet_param >= len(sheets):
            raise ValueError(
                f"indice de sheet fora do intervalo: {sheet_param} "
                f"(workbook tem {len(sheets)} planilha(s))"
            )
        return sheets[sheet_param]

    if isinstance(sheet_param, str):
        if sheet_param not in wb.sheetnames:
            raise ValueError(
                f"sheet '{sheet_param}' nao encontrada "
                f"(disponiveis: {wb.sheetnames})"
            )
        return wb[sheet_param]

    raise ValueError("'sheet' deve ser nome (string) ou indice (int)")


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return str(value)
